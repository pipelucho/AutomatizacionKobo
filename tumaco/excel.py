import json
from dataclasses import dataclass

import openpyxl
import pandas as pd

from . import config

COL_NOMBRE, COL_USUARIO, COL_CASA = 1, 2, 3  # A, B, C
COL_REALIZADO, COL_FORMULARIO_2, COL_FORMULARIO_3 = 4, 5, 6  # D, E, F
# Columna donde se marca cada formulario (numero de formulario -> columna)
COLUMNA_FORMULARIO = {1: COL_REALIZADO, 2: COL_FORMULARIO_2, 3: COL_FORMULARIO_3}
ENCABEZADOS = {COL_REALIZADO: "REALIZADO:", COL_FORMULARIO_2: "FORMULARIO 2:", COL_FORMULARIO_3: "FORMULARIO 3:"}
MARCA_REALIZADO = "Si"
MARCA_A_MANO = "Revisar a mano"  # se envio, pero algo quedo para hacerlo a mano (ej. una cedula)
MARCA_OTRO = "Tiene otro"  # la persona que atiende manifiesta ser "Otro": no se procesa
# Marcas que no se pudieron escribir porque el Excel estaba abierto: [fila, columna, valor]
PENDIENTES_PATH = config.OUTPUT_DIR / "marcas_pendientes.json"


@dataclass
class Usuario:
    fila: int
    nombre: str
    usuario: str
    casa: str
    formulario1: str = ""  # valor de la columna D
    formulario2: str = ""  # valor de la columna E
    formulario3: str = ""  # valor de la columna F

    def marca(self, numero: int) -> str:
        """Lo que dice el Excel para el formulario `numero` (vacio = falta hacerlo)."""
        return getattr(self, f"formulario{numero}", "").strip()

    def siguiente(self) -> int | None:
        """Primer formulario por hacer, o None si ya termino o una marca detiene la cadena
        (ej. 'Tiene otro' o 'No encontrado' en un formulario anterior)."""
        for numero in sorted(COLUMNA_FORMULARIO):
            marca = self.marca(numero)
            if not marca:
                return numero
            if not fue_enviado(marca):
                return None
        return None


def fue_enviado(marca: str) -> bool:
    """La marca indica que el formulario quedo enviado ('Si', o 'Revisar a mano: ...')."""
    m = marca.strip().lower()
    return m == MARCA_REALIZADO.lower() or m.startswith(MARCA_A_MANO.lower())


def letra(columna: int) -> str:
    return chr(ord("A") + columna - 1)


def _texto(valor) -> str:
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def cargar_tabla() -> pd.DataFrame:
    """Lee el Excel UNA sola vez: nombre, usuario, casa, marcas de cada formulario y la fila real de Excel."""
    crudo = pd.read_excel(config.EXCEL_PATH, header=0, dtype=object)
    while len(crudo.columns) < max(COLUMNA_FORMULARIO.values()):  # las columnas D/E/F pueden no existir todavia
        crudo[f"_vacia_{len(crudo.columns)}"] = None
    tabla = pd.DataFrame({
        "nombre": crudo.iloc[:, COL_NOMBRE - 1],
        "usuario": crudo.iloc[:, COL_USUARIO - 1],
        "casa": crudo.iloc[:, COL_CASA - 1],
        **{f"formulario{n}": crudo.iloc[:, col - 1] for n, col in COLUMNA_FORMULARIO.items()},
    }).map(_texto)
    tabla["fila"] = crudo.index + 2  # la fila 1 son los encabezados
    guardar_pendientes()  # marcas de una corrida anterior en la que el Excel estaba abierto
    return aplicar_pendientes(tabla)


def pendientes(tabla: pd.DataFrame) -> list[Usuario]:
    """Usuarios con algo por hacer, en el orden del Excel: tienen un formulario sin marca y todos
    los anteriores enviados. Una marca como 'Tiene otro' o 'No encontrado' detiene a esa persona.
    """
    usuarios = [
        Usuario(fila=int(f.fila), nombre=f.nombre, usuario=f.usuario, casa=f.casa,
                **{f"formulario{n}": getattr(f, f"formulario{n}") for n in COLUMNA_FORMULARIO})
        for f in tabla[tabla["usuario"] != ""].itertuples(index=False)
    ]
    return [u for u in usuarios if u.siguiente() is not None]


def marcar_realizado(fila: int):
    """Marca el formulario 1 (columna D) apenas Kobo confirma el envio."""
    marcar(fila, MARCA_REALIZADO)


def marcar(fila: int, valor: str, columna: int = COL_REALIZADO) -> bool:
    """Escribe `valor` en la columna indicada de la fila (con openpyxl para conservar el formato).

    Cualquier valor en D hace que el formulario 1 no se repita; en E, el formulario 2.
    Si el Excel esta abierto no se detiene la corrida: la marca queda en marcas_pendientes.json
    y se escribe en el siguiente intento (o al final, o en la proxima corrida).
    """
    marcas = _leer_pendientes()
    marcas.append([fila, columna, valor])
    _guardar_pendientes(marcas)
    return guardar_pendientes()


def guardar_pendientes(esperar: bool = False) -> bool:
    """Escribe en el Excel las marcas que no se pudieron guardar. True si ya no queda ninguna.

    Con esperar=True (al terminar la corrida) pide cerrar el Excel hasta poder guardar.
    """
    marcas = _leer_pendientes()
    if not marcas:
        return True
    while True:
        try:
            wb = openpyxl.load_workbook(config.EXCEL_PATH)
            ws = wb.active
            for fila, columna, valor in marcas:
                if not _texto(ws.cell(1, columna).value) and columna in ENCABEZADOS:
                    ws.cell(1, columna).value = ENCABEZADOS[columna]
                ws.cell(fila, columna).value = valor
            wb.save(config.EXCEL_PATH)
            PENDIENTES_PATH.unlink(missing_ok=True)
            if len(marcas) > 1:
                print(f"   Excel: se guardaron {len(marcas)} marca(s) que estaban pendientes.")
            return True
        except PermissionError:
            if not esperar:
                print(f"   [!] {config.EXCEL_PATH.name} esta abierto en Excel: la marca queda guardada en "
                      f"{PENDIENTES_PATH.name} y se escribira cuando lo cierres. Se continua.")
                return False
            input(f"Cierra {config.EXCEL_PATH.name} en Excel y presiona Enter para guardar "
                  f"{len(marcas)} marca(s) pendiente(s)...")


def aplicar_pendientes(tabla: pd.DataFrame) -> pd.DataFrame:
    """Pone en la tabla leida las marcas que aun no llegaron al Excel (para no repetir formularios)."""
    campo = {col: f"formulario{n}" for n, col in COLUMNA_FORMULARIO.items()}
    for fila, columna, valor in _leer_pendientes():
        if columna in campo:
            tabla.loc[tabla["fila"] == fila, campo[columna]] = valor
    return tabla


def _leer_pendientes() -> list:
    try:
        return json.loads(PENDIENTES_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, ValueError):
        return []


def _guardar_pendientes(marcas: list):
    PENDIENTES_PATH.parent.mkdir(parents=True, exist_ok=True)
    PENDIENTES_PATH.write_text(json.dumps(marcas), encoding="utf-8")
