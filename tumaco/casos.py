"""Registro de casos por resolver (CASOS.xlsx): no encontrados, dudosos, sin formulario 4, cedulas a mano...

Los casos salen del estado del Excel de usuarios (y de los errores de cada corrida), asi que el
registro se puede rehacer en cualquier momento. En cada actualizacion:
- un caso nuevo entra como 'Pendiente';
- uno que ya no aparece (el bot lo completo o lo corregiste) pasa a 'Resuelto' con la fecha;
- uno resuelto que vuelve a aparecer se reabre.
Las columnas 'Notas' y 'Estado' son tuyas: si pones otro estado (ej. 'Descartado'), se respeta.
"""
import json
from collections import Counter
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import config, excel

ERRORES_PATH = config.OUTPUT_DIR / "errores.json"
AVISOS_PATH = config.OUTPUT_DIR / "avisos.json"  # decisiones de nombre/casa para verificar
PENDIENTE, RESUELTO = "Pendiente", "Resuelto"
COLUMNAS = ["Estado", "Tipo", "Formulario", "Fila Excel", "Usuario", "Nombre", "Detalle", "Que hacer",
            "Detectado", "Ultima revision", "Resuelto", "Notas", "Clave"]
ANCHOS = {"Estado": 11, "Tipo": 17, "Formulario": 11, "Fila Excel": 9, "Usuario": 9, "Nombre": 30,
          "Detalle": 60, "Que hacer": 55, "Detectado": 16, "Ultima revision": 16, "Resuelto": 16,
          "Notas": 40, "Clave": 22}

QUE_HACER = {
    "No encontrado": "Buscar en Kobo por documento o nombre. Si existe con otros datos, corregir el Excel "
                     "(o el formulario) y borrar la marca de esa columna para que el bot lo repita.",
    "Dudoso": "Revisar en Kobo los registros repetidos, dejar uno solo y borrar la marca para repetir.",
    "Sin formulario 4": "Confirmar en Kobo que la persona no tiene formulario 4 (o buscarlo por documento).",
    "Revisar a mano": "Correr herramientas\\rehacer_formulario2.py (cedulas) o corregirlo en Kobo.",
    "Tiene otro": "Quien atendio la visita manifesto ser 'Otro': definir si se procesa.",
    "Vivienda inexistente (RETIRADO)": "Lote vacio: no se proceso; en la hoja queda RETIRADO (FOTOGRAFIA VIVIENDA).",
    "Hoja pendiente": "Abrir el Chrome de la hoja y correr el bot (solo llena la hoja de esa persona).",
    "Faltan datos": "Correr herramientas\\completar_datos_form1.py.",
    "Error": "Revisar el log de la corrida (salida\\logs); el bot lo reintenta en la siguiente corrida.",
    "Verificar nombre": "Kobo tiene otro nombre que el Excel y se dejo el de Kobo. Confirmar quien es; si esta bien, "
                        "cambiar el Estado a 'Verificado'.",
    "Registro ajeno": "El bot pudo editar el registro de OTRA persona con el mismo nombre (antes aceptaba un "
                      "registro solo por el nombre). Revisar ese registro en Kobo; el formulario se rehace.",
    "Registro duplicado": "Hay mas de un registro de Kobo que coincide con la persona. Confirmar cual es el "
                          "correcto y si el otro es un duplicado.",
    "Verificar casa": "Kobo, el Excel y la hoja no coinciden en la casa. Confirmar cual es la correcta; si esta "
                      "bien, cambiar el Estado a 'Verificado'.",
    "Otro": "Revisar la marca escrita en el Excel.",
}


# ---------- Deteccion ----------

def _tipo_marca(numero: int, marca: str) -> str | None:
    m = marca.strip().lower()
    if not m or excel.fue_enviado(m) and not m.startswith(excel.MARCA_A_MANO.lower()):
        return None
    if m.startswith(excel.MARCA_NO_ENCONTRADO.lower()):
        return "Sin formulario 4" if numero in excel.FORMULARIOS_OPCIONALES else "No encontrado"
    if m.startswith("dudoso"):
        return "Dudoso"
    if m.startswith(excel.MARCA_A_MANO.lower()):
        return "Revisar a mano"
    if m.startswith(excel.MARCA_OTRO.lower()):
        return "Tiene otro"
    if m.startswith(excel.MARCA_INEXISTENTE.lower()):
        return "Vivienda inexistente (RETIRADO)"
    return "Otro"


def detectar(tabla) -> list[dict]:
    """Casos abiertos segun el Excel de usuarios y los errores de las corridas."""
    errores = _leer_errores()
    avisos = _leer_json(AVISOS_PATH)
    casos = []

    def caso(u, tipo, formulario, detalle):
        casos.append({"Clave": f"{u.usuario}|{formulario}|{tipo}", "Tipo": tipo, "Formulario": formulario,
                      "Fila Excel": u.fila, "Usuario": u.usuario, "Nombre": u.nombre, "Detalle": detalle,
                      "Que hacer": QUE_HACER.get(tipo, "")})

    for u in excel.usuarios(tabla):
        for numero in sorted(excel.COLUMNA_FORMULARIO):
            marca = u.marca(numero)
            tipo = _tipo_marca(numero, marca)
            if tipo:
                caso(u, tipo, f"F{numero}", marca)
            elif not marca and (error := errores.get(f"{u.usuario}|{numero}")):
                caso(u, "Error", f"F{numero}", f"{error['fecha']}: {error['mensaje']}")
        if u.terminado() and not u.hoja:
            caso(u, "Hoja pendiente", "Hoja", "Terminó los formularios pero su fila de la hoja de Google no está llena")
        if excel.fue_enviado(u.formulario1):
            faltan = [n for n, v in (("propietario (H)", u.propietario), ("documento (I)", u.documento),
                                     ("documento anexo (J)", u.anexo)) if not v]
            if u.propietario.lower() == "sin respuesta":
                faltan.append("propietario sin respuesta en el formulario 1")
            if faltan:
                caso(u, "Faltan datos", "F1", "Falta en el Excel: " + ", ".join(faltan))
        # Decisiones de nombre/casa para verificar (quedan hasta que les cambies el Estado)
        for clave, aviso in avisos.items():
            if clave.startswith(f"{u.usuario}|"):
                caso(u, aviso["tipo"], aviso["formulario"], f"{aviso['fecha']}: {aviso['texto']}")
    return casos


# ---------- Errores de las corridas ----------

def anotar_error(usuario, numero: int, mensaje: str):
    """Guarda el ultimo error de un formulario; deja de ser caso cuando el formulario tenga marca."""
    errores = _leer_errores()
    errores[f"{usuario.usuario}|{numero}"] = {"fecha": f"{datetime.now():%Y-%m-%d %H:%M}",
                                              "mensaje": mensaje.splitlines()[0][:300]}
    _guardar_json(ERRORES_PATH, errores)


def anotar_aviso(usuario, numero: int, tipo: str, texto: str):
    """Guarda una decision para verificar (ej. 'Verificar nombre'); uno por persona, formulario y tipo."""
    avisos = _leer_json(AVISOS_PATH)
    avisos[f"{usuario.usuario}|F{numero}|{tipo}"] = {"tipo": tipo, "formulario": f"F{numero}", "texto": texto,
                                                      "fecha": f"{datetime.now():%Y-%m-%d %H:%M}"}
    _guardar_json(AVISOS_PATH, avisos)


def _leer_errores() -> dict:
    return _leer_json(ERRORES_PATH)


def _leer_json(ruta) -> dict:
    try:
        return json.loads(ruta.read_text(encoding="utf-8"))
    except (FileNotFoundError, ValueError):
        return {}


def _guardar_json(ruta, datos: dict):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")


# ---------- Registro en CASOS.xlsx ----------

def actualizar(tabla) -> list[dict] | None:
    """Actualiza CASOS.xlsx con los casos de ahora. Devuelve todos los casos (o None si el archivo esta abierto)."""
    ahora = f"{datetime.now():%Y-%m-%d %H:%M}"
    detectados = {c["Clave"]: c for c in detectar(tabla)}
    filas = _leer()
    for fila in filas:
        actual = detectados.pop(fila["Clave"], None)
        if actual:
            fila.update({k: actual[k] for k in ("Fila Excel", "Nombre", "Detalle", "Que hacer")})
            fila["Ultima revision"] = ahora
            if fila["Estado"] == RESUELTO:  # volvio a aparecer
                fila["Estado"], fila["Resuelto"] = PENDIENTE, ""
                fila["Notas"] = (fila.get("Notas") or "") + f" [reabierto {ahora}]"
        elif fila["Estado"] == PENDIENTE:  # ya no aparece: el bot lo completo o se corrigio
            fila["Estado"], fila["Resuelto"] = RESUELTO, ahora
    for nuevo in detectados.values():
        filas.append({**nuevo, "Estado": PENDIENTE, "Detectado": ahora, "Ultima revision": ahora,
                      "Resuelto": "", "Notas": ""})

    filas.sort(key=lambda f: (f["Estado"] != PENDIENTE, f["Tipo"], int(f["Fila Excel"] or 0), f["Formulario"]))
    try:
        _escribir(filas)
    except PermissionError:
        print(f"   [!] {config.CASOS_PATH.name} esta abierto: no se actualizo. Se actualiza en la proxima "
              f"corrida o con herramientas\\casos.py")
        return None
    return filas


def resumen(filas: list[dict]) -> str:
    pendientes = Counter(f["Tipo"] for f in filas if f["Estado"] == PENDIENTE)
    if not pendientes:
        return "Casos por resolver: ninguno"
    detalle = ", ".join(f"{tipo}: {n}" for tipo, n in pendientes.most_common())
    return f"Casos por resolver: {sum(pendientes.values())} ({detalle}) -> {config.CASOS_PATH.name}"


def _leer() -> list[dict]:
    if not config.CASOS_PATH.exists():
        return []
    ws = openpyxl.load_workbook(config.CASOS_PATH)["Casos"]
    encabezados = [c.value for c in ws[1]]
    filas = []
    for valores in ws.iter_rows(min_row=2, values_only=True):
        fila = {h: ("" if v is None else v) for h, v in zip(encabezados, valores) if h}
        if fila.get("Clave"):
            filas.append({c: fila.get(c, "") for c in COLUMNAS})
    return filas


def _escribir(filas: list[dict]):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Casos"
    ws.append(COLUMNAS)
    for fila in filas:
        ws.append([fila.get(c, "") for c in COLUMNAS])

    negrita, blanco = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="305496")
    colores = {PENDIENTE: PatternFill("solid", fgColor="FCE4D6"), RESUELTO: PatternFill("solid", fgColor="E2EFDA")}
    for celda in ws[1]:
        celda.font, celda.fill = negrita, blanco
    for i, nombre in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ANCHOS.get(nombre, 14)
    col_estado = COLUMNAS.index("Estado") + 1
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
        if relleno := colores.get(fila[col_estado - 1].value):
            fila[col_estado - 1].fill = relleno
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Hoja de resumen: cuantos casos hay de cada tipo y estado
    rs = wb.create_sheet("Resumen")
    estados = sorted({f["Estado"] for f in filas} | {PENDIENTE, RESUELTO}, key=lambda e: (e != PENDIENTE, e))
    rs.append(["Tipo", *estados, "Que hacer"])
    for celda in rs[1]:
        celda.font, celda.fill = negrita, blanco
    conteo = Counter((f["Tipo"], f["Estado"]) for f in filas)
    for tipo in sorted({f["Tipo"] for f in filas}):
        rs.append([tipo, *[conteo.get((tipo, e), 0) for e in estados], QUE_HACER.get(tipo, "")])
    rs.column_dimensions["A"].width = 18
    rs.column_dimensions[openpyxl.utils.get_column_letter(len(estados) + 2)].width = 90
    wb.save(config.CASOS_PATH)
