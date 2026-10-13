"""Hoja local para el procesamiento aparte (HOJA_LOCAL en .env, ej. HOJA_APARTE.xlsx).

Estos usuarios no se escriben en la hoja de Google: lo que iria en su fila (F a L) se guarda en un Excel con las
mismas columnas, y la columna K (RESPONSABLE) queda en blanco para que la llene despues la persona encargada.
Las columnas A a E (No., nombre y direccion iniciales) se copian de la hoja de Google al agregar al usuario.
"""
from openpyxl import load_workbook

from . import config, sheets
from .texto import reparar

FILA_ENCABEZADOS = 1


def _columna(letra: str) -> int:
    return ord(letra) - ord("A") + 1


def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _filas(ws) -> dict[str, int]:
    """No. de usuario (columna A) -> fila del Excel."""
    return {_texto(ws.cell(n, 1).value): n for n in range(FILA_ENCABEZADOS + 1, ws.max_row + 1)
            if _texto(ws.cell(n, 1).value)}


def referencias() -> dict[str, dict]:
    """Igual que sheets.referencias, desde la hoja local: nombre inicial (B) y direccion inicial (E)."""
    ws = load_workbook(config.HOJA_LOCAL, read_only=True).active
    refs = {}
    for numero, n in _filas(ws).items():
        nombre = _texto(ws.cell(n, 2).value)
        if nombre:
            refs[numero] = {"nombre": reparar(nombre), "fila": n,
                            "direccion": reparar(_texto(ws.cell(n, _columna(sheets.COL_DIRECCION_INICIAL)).value))}
    return refs


def actualizar_usuario(usuario, retirado: bool = False) -> bool:
    """Llena F a L de la fila del usuario en la hoja local, con las mismas reglas que la hoja de Google y la
    columna K (RESPONSABLE) en blanco. Devuelve True si quedo guardado."""
    nombre_archivo = config.HOJA_LOCAL.name
    wb = load_workbook(config.HOJA_LOCAL)
    ws = wb.active
    fila = _filas(ws).get(str(usuario.usuario))
    if not fila:
        print(f"   {nombre_archivo}: el usuario {usuario.usuario} no esta en la columna A.")
        return False
    inicial = _texto(ws.cell(fila, 2).value)
    direccion = _texto(ws.cell(fila, _columna(sheets.COL_DIRECCION_INICIAL)).value)
    if not sheets.es_de_la_persona({"nombre": reparar(inicial), "direccion": reparar(direccion)}, usuario):
        print(f"   [!] {nombre_archivo}: la fila {fila} del usuario {usuario.usuario} es de '{inicial}' ({direccion}), "
              f"no de '{usuario.nombre}' casa {usuario.casa}: revisa el No. de usuario. No se llena.")
        return False
    faltan = [n for n, v in (("documento", usuario.documento), ("propietario", usuario.propietario),
                             ("documento anexo", usuario.anexo)) if not v]
    if faltan and not retirado:
        print(f"   {nombre_archivo}: faltan datos en el Excel ({', '.join(faltan)}); no se llena la fila {fila}.")
        return False

    valores = (sheets.valores_retirado if retirado else sheets.valores_usuario)(usuario, direccion)
    valores[sheets.COL_RESPONSABLE] = ""  # la llena despues la persona encargada
    print(f"   {nombre_archivo}, fila {fila} (usuario {usuario.usuario}):")
    for col, valor in valores.items():
        ws.cell(fila, _columna(col)).value = valor or None
        print(f"     {col}{fila} = {valor or '(en blanco)'}")
    try:
        wb.save(config.HOJA_LOCAL)
    except PermissionError:
        print(f"   [!] {nombre_archivo} esta abierto: cierralo; la fila queda pendiente para la proxima corrida.")
        return False
    return True
