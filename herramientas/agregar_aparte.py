"""Agrega un usuario al procesamiento aparte (aparte.py), sin escribir nada en la hoja de Google ni en Kobo.

- Lee su fila en la hoja de Google (columnas A-E e I) y la copia a HOJA_APARTE.xlsx.
- Lo agrega a USUARIOS_APARTE.xlsx (nombre como en la cedula, usuario y casa).
- Guarda en ENTRADA_APARTE/<usuario>/ los archivos que reemplazan los de Kobo (en JPG): CEDULA_CARA.jpg,
  CEDULA_REVERSO.jpg y FACTURA.jpg.
- Muestra sus registros en Kobo (solo lectura) para revisar que existan y cuadren.

Uso (con el Chrome de la hoja abierto y los Excel cerrados):
  .venv\\Scripts\\python herramientas\\agregar_aparte.py --usuario 123 --casa 45 --nombre "NOMBRE COMO EN LA CEDULA"
      --cara cedula.webp --reverso reverso.webp --factura factura.webp [--registro "F1=...; F2=..."]
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook  # noqa: E402
from PIL import Image  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

from tumaco import analisis, config, sheets  # noqa: E402
from tumaco.browser import lanzar  # noqa: E402
from tumaco.texto import digitos, reparar, similitud_nombre  # noqa: E402

USUARIOS = config.ROOT / "USUARIOS_APARTE.xlsx"
HOJA = config.ROOT / "HOJA_APARTE.xlsx"
ENTRADA = config.ROOT / "ENTRADA_APARTE"
# Por formulario: nombre, documento, No. usuario, direccion, barrio
CAMPOS = {
    1: ("grp_datos_persona_atendida/nombre_persona_atendida", "grp_datos_persona_atendida/documento_persona_atendida",
        "grp_datos_persona_atendida/no_usuario", "grp_datos_persona_atendida/direccion", "grp_datos_persona_atendida/barrio"),
    2: ("grp_declaracion/nombre_usuario_decl", "grp_declaracion/documento_usuario_decl", "grp_declaracion/No_Usuario",
        "grp_declaracion/direccion_usuario_decl", "grp_declaracion/barrio_usuario_decl"),
    3: ("datos_usuario/nombre", "datos_usuario/documento", "datos_usuario/no_usuario", "datos_usuario/direccion",
        "datos_usuario/barrio"),
    4: ("grp_usuario_inmueble/nombre_usuario", "grp_usuario_inmueble/documento_usuario", "grp_usuario_inmueble/no_usuario",
        "grp_usuario_inmueble/direccion", "grp_usuario_inmueble/barrio"),
}


def registros_kobo(usuario: str, casa: str, nombre: str, documento: str):
    """Registros de Kobo que parecen de esta persona: mismo No. usuario, documento, nombre o casa."""
    with sync_playwright() as p:
        ctx = lanzar(p, headless=True)
        try:
            req = ctx.request
            activos = req.get(f"{config.KOBO_URL}/api/v2/assets/", params={"format": "json", "q": "asset_type:survey",
                                                                            "limit": "100"}).json()["results"]
            for n, titulo in enumerate(config.FORMULARIOS, start=1):
                uid = next(a["uid"] for a in activos if analisis.norm(a["name"]).startswith(analisis.norm(titulo)[:20]))
                datos = req.get(f"{config.KOBO_URL}/api/v2/assets/{uid}/data/",
                                params={"format": "json", "limit": "30000"}).json()["results"]
                nom, doc, us, dire, barrio = CAMPOS[n]
                hallados = 0
                for r in datos:
                    v = lambda k: " ".join(str(r.get(k, "") or "").split())  # noqa: E731
                    if (digitos(v(us)) == usuario or (documento and documento in re.sub(r"\D", "", v(doc)))
                            or similitud_nombre(nombre, v(nom)) >= 0.85 or casa in digitos(v(dire), todos=True)
                            and similitud_nombre(nombre, v(nom)) >= 0.6):
                        hallados += 1
                        anexos = {k.split("/")[-1]: r[k] for k in r
                                  if any(x in k for x in ("soporte_propiedad", "cedula_usuario"))}
                        print(f"   F{n} _id {r['_id']} | {v(nom)} | doc {v(doc)} | usuario {v(us)} | {v(dire)} | "
                              f"{v(barrio)} | {anexos or ''}")
                if not hallados:
                    print(f"   F{n}: no aparece")
        finally:
            ctx.close()


def main():
    a = argparse.ArgumentParser(description="Agregar un usuario al procesamiento aparte")
    a.add_argument("--usuario", required=True)
    a.add_argument("--casa", required=True)
    a.add_argument("--nombre", required=True, help="como en la cedula")
    a.add_argument("--documento", default="", help="numero de la cedula (para buscarlo en Kobo)")
    a.add_argument("--cara")
    a.add_argument("--reverso")
    a.add_argument("--factura")
    a.add_argument("--registro", default="", help="registros de Kobo ya decididos, ej. 'F1=123456789; F2=123456790' "
                                                     "(cuando otra persona tiene el mismo No. de usuario)")
    args = a.parse_args()
    usuario, casa, nombre = args.usuario.strip(), args.casa.strip(), " ".join(args.nombre.upper().split())

    # Hoja de Google (solo lectura)
    with sheets.conectar_hoja() as h:
        filas = sheets.leer_filas(h)
    n = sheets.buscar_fila(filas, usuario)
    if not n:
        sys.exit(f"El usuario {usuario} no esta en la hoja de Google: no se agrega.")
    fila = [reparar(sheets._celda(filas, n, c)) for c in "ABCDEFGHIJKL"]  # la hoja trae letras danadas
    print(f"Hoja de Google, fila {n}: {fila[1]} | {fila[4]} | {fila[8]}")
    ref = {"nombre": fila[1], "direccion": fila[4]}
    falso = argparse.Namespace(nombre=nombre, nombre_final="", casa=casa, casa_final="", usuario=usuario)
    if not sheets.es_de_la_persona(ref, falso):
        sys.exit(f"La fila {n} de la hoja ('{fila[1]}', {fila[4]}) no parece de {nombre} casa {casa}: no se agrega.")

    # USUARIOS_APARTE.xlsx
    wb = load_workbook(USUARIOS)
    ws = wb.active
    if any(str(ws.cell(r, 2).value).strip() == usuario for r in range(2, ws.max_row + 1)):
        print(f"{USUARIOS.name}: el {usuario} ya estaba")
    else:
        ws.append([nombre, int(usuario), int(casa) if casa.isdigit() else casa])
        if args.registro:
            ws.cell(ws.max_row, 15).value = args.registro  # columna O (REGISTRO KOBO)
        wb.save(USUARIOS)
        print(f"{USUARIOS.name}: agregado en la fila {ws.max_row}")

    # HOJA_APARTE.xlsx: A-E e I de la hoja de Google
    hb = load_workbook(HOJA)
    hs = hb.active
    if any(str(hs.cell(r, 1).value).strip() == usuario for r in range(2, hs.max_row + 1)):
        print(f"{HOJA.name}: el {usuario} ya estaba")
    else:
        hs.append([int(usuario)] + [x or None for x in fila[1:5]] + [None] * 3 + [fila[8] or None])
        hb.save(HOJA)
        print(f"{HOJA.name}: agregado en la fila {hs.max_row}")

    # Archivos de entrada
    carpeta = ENTRADA / usuario
    carpeta.mkdir(parents=True, exist_ok=True)
    for origen, destino in ((args.cara, "CEDULA_CARA.jpg"), (args.reverso, "CEDULA_REVERSO.jpg"),
                            (args.factura, "FACTURA.jpg")):
        if origen:
            img = Image.open(origen).convert("RGB")
            img.save(carpeta / destino, quality=92)
            print(f"{carpeta.relative_to(config.ROOT) / destino}: {img.size[0]}x{img.size[1]}")

    print("\nRegistros en Kobo (solo lectura):")
    registros_kobo(usuario, casa, nombre, re.sub(r"\D", "", args.documento))


if __name__ == "__main__":
    main()
